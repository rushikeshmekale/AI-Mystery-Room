from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out = Path(r'C:\Users\rushi\Documents\Codex\2026-10-26\i-x20\outputs\ai_mystery_room_spec.docx')
doc = Document()
sec = doc.sections[0]
sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.5); styles['Normal'].font.color.rgb=RGBColor(44,54,66)
styles['Normal'].paragraph_format.space_after=Pt(5)
for nm,sz,col in [('Title',30,'172B4D'),('Heading 1',19,'172B4D'),('Heading 2',13,'176B87'),('Heading 3',10.5,'176B87')]:
 s=styles[nm]; s.font.name='Aptos Display'; s.font.size=Pt(sz); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(col)
 s.paragraph_format.space_before=Pt(10 if nm!='Title' else 0); s.paragraph_format.space_after=Pt(5)

# header/footer
header=sec.header.paragraphs[0]; header.text='AI MYSTERY ROOM  /  PRODUCT AND ENGINEERING SPECIFICATION'; header.style='Caption'; header.runs[0].font.color.rgb=RGBColor(100,113,128); header.runs[0].font.size=Pt(8)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run('Single source of truth  •  Version 1.0  •  26 October 2026'); run.font.size=Pt(8); run.font.color.rgb=RGBColor(100,113,128)

def shade(cell, fill):
 tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def table(headers, rows, widths=None):
 t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Light Shading Accent 1'
 for i,h in enumerate(headers):
  c=t.rows[0].cells[i]; c.text=h; shade(c,'176B87')
  for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(8.5)
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):
   cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for p in cells[i].paragraphs:
    p.paragraph_format.space_after=Pt(2)
    for r in p.runs: r.font.size=Pt(8.5)
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 doc.add_paragraph().paragraph_format.space_after=Pt(1)
 return t
def p(text=''): doc.add_paragraph(text)
def h(text,level=1): doc.add_heading(text,level=level)
def bullets(items):
 for x in items: doc.add_paragraph(x,style='List Bullet')
def nums(items):
 for x in items: doc.add_paragraph(x,style='List Number')

# cover
t=doc.add_paragraph(style='Title'); t.add_run('AI Mystery Room')
para=doc.add_paragraph(); r=para.add_run('Product and Engineering Specification'); r.bold=True; r.font.size=Pt(16); r.font.color.rgb=RGBColor(23,107,135)
p('A build-ready source of truth for a browser-based, AI-generated investigation game.')
table(['Document status','Version','Last updated','Product owner'],[['MVP specification','1.0','26 October 2026','Project team']])
h('Purpose',1)
p('This specification defines the user experience, system behavior, data contracts, AI generation rules, and acceptance criteria for AI Mystery Room. Codex should treat it as the authoritative product and implementation brief. If implementation needs to depart from a requirement, update this document first and record the decision in the change log.')
h('At a glance',2)
table(['Area','Decision'],[
['Product','A user supplies a short fictional premise. The service turns it into a solvable mystery, presents evidence, and asks the user to identify the culprit.'],
['Primary stack','MongoDB, Express, React, Node.js (MERN).'],
['AI provider','Google Gemini, called only by the server.'],
['MVP format','Single-player, one case at a time, responsive web app; anonymous sessions are supported.'],
['Core principle','One canonical case solution drives every clue, character statement, reveal, and evaluation.'],
['Media','Text and generated/illustrative evidence in MVP. External image/video generation is optional and must degrade gracefully.']])

h('1. Product overview')
h('1.1 Problem and opportunity',2)
p('People enjoy interactive mysteries but often want a story tailored to a theme, setting, or idea they already have. AI Mystery Room lets them contribute a premise and play a short investigation assembled around it. The fun comes from examining varied evidence, forming a theory, and receiving a fair explanation.')
h('1.2 Product vision',2)
p('Create a lightweight mystery game in which every case feels specific to its premise while remaining internally consistent, safe for general audiences, and solvable from the evidence shown.')
h('1.3 Goals',2)
bullets(['Let a user create a case from a short premise without needing a writing prompt or account.','Produce a coherent fictional incident, a cast of suspects, evidence, and one authoritative culprit.','Let players inspect evidence, record a theory, submit one final accusation, and see a spoiler-aware explanation.','Keep generation, storage, and evaluation reliable enough that a failed AI call never strands a user.','Make the experience work on current desktop and mobile browsers.'])
h('1.4 Non-goals for MVP',2)
bullets(['Real crime solving, law enforcement advice, or use of real people, real private messages, or real surveillance footage.','Multiplayer, live collaboration, public case sharing, leaderboards, payments, or a creator marketplace.','Long-form narrative campaigns, branching endings, voice interaction, or unrestricted chat with suspects.','Guaranteeing a true-to-life forensic process or treating generated evidence as authentic records.'])
h('1.5 Target users and use cases',2)
table(['User','Need','MVP response'],[
['Curious player','A quick, self-contained mystery to solve.','Enter a premise, inspect a case, accuse a suspect, read the reveal.'],
['Returning player','Resume or review an in-progress case.','Use a locally stored session token to reopen the case on the same device.'],
['Builder / operator','Diagnose invalid output and service failures.','Structured logs, validation, retry limits, and safe error states; no public admin console required.']])

h('2. Product requirements')
h('2.1 Core gameplay loop',2)
nums(['User enters a premise and selects optional tone, difficulty, and evidence style.','Server validates the request and generates a complete case blueprint.','Server validates internal consistency and persists the playable case before returning it.','User reviews the case brief, suspects, timeline, and evidence.','User may open evidence in any order and optionally save private notes.','User selects one suspect, optionally enters a short rationale, and submits a final accusation.','Server evaluates the selection against the stored answer and returns the case resolution, explanation, and score summary.'])
h('2.2 Functional requirements',2)
table(['ID','Requirement','Priority'],[
['FR-01','Accept a premise between 20 and 600 characters after trimming. Show a live character counter and clear validation.','P0'],
['FR-02','Offer optional tone (cozy, suspenseful, noir, light sci-fi) and difficulty (easy, standard, challenging). Defaults: suspenseful and standard.','P1'],
['FR-03','Generate a case with a title, premise adaptation, incident, victim/target or affected party, 3–5 suspects, timeline, 6–10 evidence items, and one culprit.','P0'],
['FR-04','Display case overview, suspect profiles, timeline, and evidence in a clear investigation workspace.','P0'],
['FR-05','Allow evidence to be opened, revisited, and marked reviewed. Do not reveal the culprit before accusation.','P0'],
['FR-06','Allow optional notes stored against the case session. Notes are private to the session and are not sent to Gemini.','P1'],
['FR-07','Accept one final accusation. Confirm before submission and prevent duplicate submission.','P0'],
['FR-08','Return correct/incorrect outcome and a spoiler-marked explanation citing evidence IDs.','P0'],
['FR-09','Persist case, evidence, progress, notes, and outcome in MongoDB. A session token permits anonymous resume.','P0'],
['FR-10','Handle generation errors, timeouts, rate limits, and malformed model output with recoverable UI and no partial playable case.','P0'],
['FR-11','Support replay by creating a new case. Replaying the same case or changing its solution is out of scope.','P1'],
['FR-12','Provide accessible controls, keyboard navigation, responsive layout, loading feedback, and reduced-motion behavior.','P0']])
h('2.3 User stories',2)
bullets(['As a player, I can describe a story idea and receive a playable mystery so the case feels personal.','As a player, I can inspect every clue before accusing so I can make an informed choice.','As a player, I can revisit evidence and keep notes so I can track contradictions.','As a player, I can submit an accusation once and see why the answer is supported.','As a player, I can resume on the same browser after refresh without creating an account.'])
h('2.4 User flow and screens',2)
table(['Screen / state','Required content and behavior'],[
['Home / create case','Product name and concise explanation; premise textarea; optional tone/difficulty; create button; validation and privacy note that stories should be fictional.'],
['Generating','Progress indication with honest wording (case preparation); prevent duplicate requests; allow safe cancel if supported. Do not imply exact AI step completion.'],
['Investigation','Case title and brief; incident details; suspect roster; timeline; evidence index with type and reviewed state; evidence detail panel/dialog; notes area; accusation control.'],
['Accusation confirmation','Selected suspect and clear irreversible-submission notice; confirm or return to investigation.'],
['Resolution','Correct/incorrect result; culprit identity; reasoning chain; each decisive clue linked by ID/title; explanation of misleading clues; score summary; create another case.'],
['Error / unavailable','Plain-language message, retry path, preserved premise where possible, and no fake case or invented success.']])

h('3. Mystery and evidence design')
h('3.1 Case blueprint contract',2)
p('Gemini generates a structured case blueprint, never a free-form narrative that the frontend must parse. The server validates it before persistence. The stored answer key is server-only and excluded from all investigation responses.')
table(['Field','Constraints'],[
['caseTitle','3–80 characters; suitable for general audiences.'],
['caseSummary','80–500 characters; sets stakes but does not disclose the answer.'],
['incident','Clear fictional incident and time/place; no graphic detail.'],
['suspects','3–5 entries; unique stable IDs, names, roles, motives, alibis, and public profiles.'],
['timeline','4–10 timestamped events; events must not contradict established evidence.'],
['evidence','6–10 items; each has unique ID, title, type, content, source, timestamp if relevant, and related suspect IDs.'],
['solution','Exactly one culprit ID; motive, opportunity, method, and a concise reasoning chain. Server-only.'],
['clueMap','2–4 decisive clue references, 1–3 red herrings with fair explanations, plus evidence-to-claim links. Server-only until reveal.'],
['assets','Optional asset descriptors, alt text, and generation prompt metadata; never embed unbounded binary data in the case document.']])
h('3.2 Fair-play rules',2)
bullets(['The culprit must be among the displayed suspects and the stored answer must use a valid suspect ID.','At least two independent evidence items must support the culprit; at least one should distinguish the culprit from the strongest alternative.','The evidence and timeline must support the solution without requiring outside knowledge or guessing the author’s intent.','Each suspect receives plausible motive and opportunity details; motive alone is never sufficient proof.','Every red herring must be explainable in the resolution and must not falsely establish guilt.','The case brief and evidence shown before accusation must not name the culprit as the answer or expose the hidden solution fields.','If consistency validation fails, regenerate once using validation feedback; if it still fails, return a recoverable generation error. Do not serve an unvalidated case.'])
h('3.3 Evidence types and presentation',2)
table(['Type','MVP representation','Specific requirement'],[
['Message thread','Fictional chat transcript styled as a generic messaging interface.','Label clearly as fictional; do not copy WhatsApp branding, logo, or real product UI. Use invented participants and content.'],
['Image','Illustrative image, optional in MVP.','If image generation is unavailable, show a designed evidence card with a text description. Include alt text and source label.'],
['Camera / recording','Text transcript, event log, or storyboard description; optional generated clip in a future release.','Do not imply real surveillance footage. A transcript must identify gaps/limitations and avoid claiming certainty beyond the content.'],
['Document / receipt','Fictional report, note, schedule, or receipt rendered as an accessible card.','Avoid real addresses, phone numbers, account identifiers, or personal data.'],
['Witness statement','Short attributed fictional account.','Statements may be unreliable but must be reconcilable with the case timeline and evidence.']])
p('For MVP, prioritize readable, text-first evidence. Binary assets are an enhancement, not a dependency for creating or solving a case. Any generated visual or transcript must be conspicuously identified as fictional/synthetic in the interface.')
h('3.4 Difficulty and scoring',2)
table(['Difficulty','Design intent','Scoring suggestion'],[
['Easy','6–7 clues, direct corroboration, few red herrings.','Correct accusation: 100; incorrect: 0.'],
['Standard','7–9 clues, several leads, one or two fair red herrings.','Correct accusation: 100; optional rationale feedback is qualitative.'],
['Challenging','8–10 clues, subtler contradictions, stronger alternative suspect.','Correct accusation: 100; optional hint use can be recorded for future scoring.']])
p('MVP scoring is intentionally transparent and outcome-based. Do not use an LLM to decide whether an accusation is correct. A rationale may be collected for reflection, but automated free-text grading is out of scope.')

h('4. UX, content, and accessibility')
h('4.1 Interaction principles',2)
bullets(['Make evidence the primary workspace and keep navigation between overview, suspects, timeline, and evidence predictable.','Keep the answer hidden until submission; label spoiler content and collapse it by default after resolution.','Show clear progress, reviewed evidence count, and selected suspect without pressuring the player.','Use fictional branding and calm suspense. Avoid graphic crime imagery and sensational treatment.','Never make a generated image the only way to understand a clue.'])
h('4.2 Accessibility',2)
bullets(['Target WCAG 2.2 AA for contrast, keyboard operation, visible focus, and semantic structure.','Use headings, labels, buttons, and dialog semantics; announce validation and generation status to assistive technology.','Provide meaningful alt text for images; transcripts must be available for any future audio/video.','Do not communicate evidence type, correctness, or review state using color alone.','Respect reduced-motion preferences; ensure zoom and small-screen layouts remain usable.'])
h('4.3 Responsive layout',2)
p('Desktop: investigation navigation/sidebar plus a central evidence panel and optional detail rail. Tablet and mobile: stack content into navigable sections; evidence detail opens as a full-width view; keep the submit accusation action reachable without obscuring content. Avoid hover-only functionality.')

h('5. Technical architecture')
h('5.1 Stack and boundaries',2)
table(['Layer','Technology / responsibility'],[
['Frontend','React SPA; responsive accessible UI; route-level case views; API client; local session token storage.'],
['Backend','Node.js + Express REST API; request validation; authentication by opaque anonymous session token; gameplay rules; Gemini orchestration.'],
['Database','MongoDB; cases and session progress. Use schema validation and indexes; never send answer key fields to public endpoints.'],
['AI','Google Gemini API through server-side adapter; JSON-constrained generation where supported; schema validation and bounded retry.'],
['Asset storage','Optional object storage for generated image/video; store URLs and metadata, not large binaries in MongoDB. MVP must function without it.'],
['Deployment','Environment-based configuration; separate frontend and API deployment is acceptable. HTTPS required in production.']])
h('5.2 Suggested repository layout',2)
p('Keep frontend and backend clearly separated in a monorepo unless the existing repository establishes another convention. Suggested layout:')
for x in ['client/  React application, pages, components, API client, styles','server/  Express app, routes, controllers, services, middleware, models','shared/  JSON schemas and shared enums only if repository tooling supports it','docs/  this specification as the authoritative requirements source']: doc.add_paragraph(x,style='List Bullet')
p('Adapt to the repository’s existing conventions where present. Do not introduce TypeScript, UI frameworks, state libraries, or infrastructure services without a documented need or an update to this spec.')
h('5.3 API contract',2)
table(['Method and path','Purpose','Request / response summary'],[
['POST /api/v1/cases','Generate and create a case.','Body: premise, tone?, difficulty?. Response: caseId, sessionToken, public case payload, generation status. Never include solution or clueMap.'],
['GET /api/v1/cases/:caseId','Fetch/resume case.','Requires session token. Returns public case, evidence, reviewed IDs, notes, status; solution excluded unless resolved.'],
['PATCH /api/v1/cases/:caseId/progress','Save progress.','Body: reviewedEvidenceIds (validated subset) and notes (bounded length). No answer-key mutation.'],
['POST /api/v1/cases/:caseId/accusation','Submit final accusation.','Body: suspectId, rationale?; requires confirmation in UI and session token. Response: outcome and reveal payload. Idempotent return of existing resolution.'],
['GET /api/v1/health','Liveness/readiness.','Health state without secrets or detailed infrastructure diagnostics.']])
p('All API responses use a consistent envelope: { data, error, requestId }. Errors include a stable code and safe user-facing message; stack traces and provider details stay server-side. Apply request size limits, rate limits to case creation and accusation endpoints, and strict CORS configuration.')
h('5.4 Data model',2)
table(['Collection','Key fields and rules'],[
['Case','_id, ownerSessionId, premise, options, title, summary, incident, suspects[], timeline[], evidence[], status, createdAt, updatedAt, expiresAt. Keep answerKey in a separately access-controlled field/object.'],
['Answer key','caseId, culpritSuspectId, motive, method, reasoningSteps[], decisiveEvidenceIds[], redHerringExplanations[]. Never serialize through public case DTO.'],
['Session','_id, tokenHash, createdAt, lastSeenAt, expiresAt. Store only a hash of the opaque bearer token where practical.'],
['Case progress','caseId, sessionId, reviewedEvidenceIds[], notes, accusation { suspectId, rationale, submittedAt }, resolution { isCorrect, resolvedAt }. Enforce one final accusation.']])
p('Recommended indexes: unique session token hash; case owner/session plus updatedAt; TTL on expired sessions and cases if retention policy is enabled. Use MongoDB schema validation or equivalent application-level validators. Do not duplicate mutable answer fields in client state.')
h('5.5 Case generation pipeline',2)
nums(['Validate input length, options, content safety, and rate limit before calling Gemini.','Construct a server-owned prompt that requests a schema-conforming fictional case and includes fair-play rules. User premise is untrusted data, not instructions.','Request structured JSON; parse strictly and validate required fields, IDs, references, lengths, count ranges, and exactly one culprit.','Run deterministic consistency checks: unique IDs; valid suspect/evidence references; decisive clues exist; clue claims do not contradict timestamps; public fields do not reveal solution.','If invalid, make at most one repair/regeneration call with concise validator feedback. Apply timeout and total token/cost limits.','Persist the validated blueprint and answer key atomically; return only a public DTO. If persistence fails, return an error and do not expose a partial case.','Log request ID, model identifier, duration, validation result, and error category. Do not log API keys, session tokens, or full user premises by default.'])
h('5.6 Gemini integration',2)
bullets(['Keep the Gemini API key on the server in an environment secret; never put it in frontend variables or source control.','Hide provider-specific code behind a generation service so model names and API changes do not leak into routes or UI.','Use a pinned/configured model identifier, bounded request timeout, token ceilings, and one retry at most for invalid structure.','Treat all model output as untrusted input. Validate and encode it before rendering; never render raw HTML from model output.','Return predictable errors for timeout, quota/rate-limit, safety block, invalid JSON, and internal failure. Preserve the user’s premise in the UI for retry.','Do not stream the answer key to the browser during generation. Do not ask Gemini to score correctness after accusation.'])
h('5.7 Session and state flow',2)
p('On case creation, issue a cryptographically random opaque session token and associate it with the case. Store the token in secure storage suitable for an anonymous app; prefer an HttpOnly, Secure, SameSite cookie if frontend and API deployment allow it. If a bearer token is used by the SPA, store it only in sessionStorage, never in a URL or persistent localStorage, and send it only over HTTPS. On resume, verify token ownership and case status. Expire anonymous sessions after a configurable inactivity period, default 30 days.')
h('5.8 Security and privacy',2)
bullets(['Use server-side authentication/ownership checks for every case and progress endpoint; random case IDs are not authorization.','Validate and bound every input; escape user/model strings; prevent injection, stored XSS, NoSQL injection, and unsafe file uploads.','Keep secrets out of responses, logs, client bundles, and error messages. Configure security headers and production HTTPS.','Minimize personal data. Warn users to use fictional premises and not submit real names, private conversations, contact details, or sensitive personal information.','Do not use submitted premises or generated cases for model training by the application. Retain data only as needed for play and configurable expiry.','Define a safe, non-graphic content policy. Block or redirect real-world allegations, identifiable private persons, sexual content involving minors, explicit violence, instructions for wrongdoing, and other disallowed content.','Generated messages, images, and recordings must be labeled fictional. Do not use real phone numbers, actual contact handles, recognizable private individuals, or fabricated claims presented as authentic evidence.'])
h('5.9 Failure and fallback behavior',2)
table(['Failure','Required behavior'],[
['Gemini timeout or provider unavailable','Show retryable message; do not create empty case; avoid repeated uncontrolled calls.'],
['Invalid or inconsistent model output','One bounded repair attempt; then fail safely and record validator category.'],
['Database unavailable','Return service unavailable; do not show generation as complete until persistence succeeds.'],
['Expired/invalid session','Offer to create a new case; do not disclose whether another user’s case ID exists.'],
['Duplicate accusation request','Return the existing resolution; never change the submitted suspect or score.'],
['Optional asset generation fails','Use the text evidence representation and continue the case.']])

h('6. Non-functional requirements')
table(['Area','Requirement'],[
['Performance','Non-AI API reads should target p95 under 500 ms under normal MVP load. Case generation may take longer; provide progress feedback and use a defined server timeout (suggested 45 seconds).'],
['Reliability','Generation must be all-or-nothing. Persisted cases must pass schema and consistency validation. Accusation submission is idempotent.'],
['Scalability','Stateless API instances; database-backed session/progress; configurable limits. Avoid designing for unsubstantiated high-scale assumptions.'],
['Compatibility','Current stable Chrome, Edge, Firefox, and Safari desktop/mobile versions at implementation time.'],
['Maintainability','Separate route/controller, domain validation, persistence, and Gemini provider adapter. Use explicit schemas and shared API types where feasible.'],
['Observability','Structured logs with request IDs, latency, provider error class, validation outcome, and response status; redact tokens and premise text by default.'],
['Accessibility','WCAG 2.2 AA target; keyboard and screen-reader usability.'],
['Cost control','Rate limit generation; cap prompt/output size and retries; configure model and usage limits via server environment.']])

h('7. Acceptance criteria')
p('MVP is complete when all P0 criteria below can be demonstrated in a deployed or locally runnable application.')
for x in ['A user can submit a valid premise and receive a complete playable case containing 3–5 suspects and 6–10 evidence items.','Invalid, empty, too-short, and too-long premises produce clear inline errors and do not call the generation service.','The public case API and frontend never expose the culprit, answer key, decisive clue map, or solution reasoning before final accusation.','All evidence references resolve; the culprit is a displayed suspect; at least two independent clues support the stored solution.','A user can browse evidence, mark it reviewed, add notes, refresh, and resume the same case with the same browser session.','The accusation flow requires a suspect selection and confirmation; duplicate requests do not change the final result.','Correct and incorrect accusations both produce a resolution with rationale and cited evidence.','Provider timeout, quota failure, malformed output, and database failure show recoverable error states without broken/partial cases.','The app is usable by keyboard, has visible focus, semantic labels, responsive mobile layout, and text alternatives for visual evidence.','Gemini credentials remain server-side, and the repo includes an environment example without real secrets.','README explains setup, required environment variables, how to run client/server, and the generation/data flow.','The application makes clear that all cases and evidence are fictional and AI-generated.']: doc.add_paragraph(x,style='List Bullet')
h('7.1 Definition of done for each feature',2)
bullets(['Requirement behavior implemented and linked to the relevant acceptance criterion.','Loading, empty, error, and success states are present.','Authorization and input validation are enforced on the server.','No secrets or answer-key data are exposed in client payloads.','Documentation and environment examples are updated when setup or configuration changes.'])

h('8. Delivery plan')
table(['Phase','Scope','Exit condition'],[
['1. Foundation','Repository structure, React shell, Express API, MongoDB connection, configuration, health route.','App starts locally; secrets are environment-based; health route works.'],
['2. Case generation','Gemini adapter, prompt/schema, validation and consistency checks, persistence, error handling.','Validated case can be generated and public DTO omits solution.'],
['3. Investigation experience','Case overview, suspects, timeline, evidence browser, progress and notes.','Player can inspect a persisted case and resume.'],
['4. Accusation and reveal','Confirmation, one-time submission, deterministic answer checking, cited explanation.','Correct and incorrect paths both resolve from stored answer key.'],
['5. Hardening','Accessibility, responsive behavior, rate limits, redaction, failure states, setup docs.','Acceptance criteria met and deployable with documented environment.']])
h('9. Decisions and assumptions')
table(['Topic','MVP decision'],[
['Accounts','No sign-up required; anonymous browser session. Add accounts only if cross-device persistence becomes a requirement.'],
['Language','English-only interface and generation for MVP.'],
['Generation timing','Synchronous create request with loading UI and a bounded timeout. Move to queued jobs only if real latency/scale justifies it.'],
['Media','Evidence can be compelling in text-first form. Image/video generation is optional and must not block play.'],
['Case retention','Configurable expiry; default 30 days from last activity. Retention and deletion behavior must be disclosed in the UI.'],
['Scoring','Binary correctness for MVP. No subjective LLM grading or hidden scoring formula.'],
['Themes','Limited preset tones; no user-supplied style imitation or branded evidence interfaces.']])
p('These are explicit implementation assumptions inferred to make the initial build bounded. Change this section and relevant requirements before implementation if product direction changes.')
h('10. Open decisions for later versions')
bullets(['Should cases be shareable by link, and what consent/moderation controls would that require?','Should accounts be added for cross-device case history?','Should image or video generation be supported, and which provider/storage/budget constraints apply?','Should the product offer age controls, content settings, or case reporting?','Should creator-authored cases and multiplayer be separate product modes?'])
h('11. Change control')
p('This document is the single source of truth for product behavior and MVP technical decisions. Update it when requirements change, with a version/date and a concise entry below. Code comments and README may explain implementation details but must not silently override requirements here.')
table(['Version','Date','Change'],[['1.0','26 October 2026','Initial product and engineering specification for the MVP.']])

# Keep table rows together and prevent awkward table splitting where supported
for t in doc.tables:
 for row in t.rows:
  trPr=row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
doc.core_properties.title='AI Mystery Room Product and Engineering Specification'
doc.core_properties.subject='Single source of truth for AI Mystery Room spec-driven development'
doc.core_properties.author='AI Mystery Room project team'
doc.save(out)
print(out)

